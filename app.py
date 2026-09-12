from __future__ import annotations

from io import BytesIO

import pandas as pd
import streamlit as st

try:
    from docx import Document
except ModuleNotFoundError:
    Document = None

from quality_workflow import QualityWorkflow


WORKFLOW_CACHE_VERSION = "unique-sop-context-v5"


@st.cache_resource
def load_workflow(cache_version: str) -> QualityWorkflow:
    return QualityWorkflow()


workflow = load_workflow(WORKFLOW_CACHE_VERSION)

st.set_page_config(page_title="Quality CAPA Review Dashboard", layout="wide")

st.markdown(
    """
    <style>
    .stMarkdown div[data-testid="stMarkdownContainer"] p,
    .stMarkdown div[data-testid="stMarkdownContainer"] li,
    .stMarkdown div[data-testid="stMarkdownContainer"] div,
    .stMarkdown div[data-testid="stMarkdownContainer"] span,
    .stExpander {
        font-size: 0.75rem !important;
        line-height: 1.45 !important;
    }
    .stExpander summary {
        font-size: 0.95rem !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

if "status_filter" not in st.session_state:
    st.session_state["status_filter"] = "non_approved"
if "severity_filter" not in st.session_state:
    st.session_state["severity_filter"] = "All"


def refresh_queue() -> None:
    status_filter = st.session_state.get("status_filter", "non_approved")
    severity_filter = st.session_state.get("severity_filter", "All")

    rows = []
    for _, row in workflow.ticket_df.iterrows():
        row_status_raw = row.get("status")
        row_status = str(row_status_raw).strip().lower() if pd.notna(row_status_raw) else None
        if status_filter == "non_approved":
            if row_status not in {None, "rejected"}:
                continue
        elif status_filter == "approved":
            if row_status != "approved":
                continue
        elif status_filter == "all":
            pass

        processed = workflow.process_ticket(row.to_dict())
        processed["status"] = row_status
        if severity_filter != "All" and processed["severity"].lower() != severity_filter.lower():
            continue
        rows.append(processed)

    st.session_state["tickets"] = rows
    if rows:
        selected_ids = [item["ticket"]["ticket_id"] for item in rows]
        if st.session_state.get("selected_ticket_id") not in selected_ids:
            st.session_state["selected_ticket_id"] = rows[0]["ticket"]["ticket_id"]
    else:
        st.session_state["selected_ticket_id"] = None


def get_ticket_by_id(ticket_id: str):
    for ticket_state in st.session_state["tickets"]:
        if ticket_state["ticket"]["ticket_id"] == ticket_id:
            return ticket_state
    return None


def persist_ticket(updated_ticket):
    ticket_id = updated_ticket["ticket"]["ticket_id"]
    if updated_ticket.get("status") in {"resolved", "approved"}:
        workflow.update_ticket_status(ticket_id, "approved")
    elif updated_ticket.get("status") == "rejected":
        workflow.update_ticket_status(ticket_id, "rejected")
    else:
        workflow.update_ticket_status(ticket_id, None)

    for idx, ticket_state in enumerate(st.session_state["tickets"]):
        if ticket_state["ticket"]["ticket_id"] == ticket_id:
            st.session_state["tickets"][idx] = updated_ticket
            break
    refresh_queue()


def build_capa_document(ticket_state) -> bytes:
    if Document is None:
        ticket = ticket_state["ticket"]
        text = (
            f"CAPA Document - {ticket['ticket_id']}\n\n"
            f"Date: {ticket.get('date', '')}\n"
            f"Product line: {ticket.get('product_line', 'General')}\n"
            f"Severity: {ticket_state['severity']}\n"
            f"Category: {ticket_state['category']}\n\n"
            f"Incident Description\n{ticket.get('description', '')}\n\n"
            f"CAPA Recommendation\n{ticket_state.get('final_output') or ticket_state['draft']}\n"
        )
        return text.encode("utf-8")

    document = Document()
    ticket = ticket_state["ticket"]
    document.add_heading(f"CAPA Document - {ticket['ticket_id']}", level=1)
    document.add_paragraph(f"Date: {ticket.get('date', '')}")
    document.add_paragraph(f"Product line: {ticket.get('product_line', 'General')}")
    document.add_paragraph(f"Severity: {ticket_state['severity']}")
    document.add_paragraph(f"Category: {ticket_state['category']}")
    document.add_heading("Incident Description", level=2)
    document.add_paragraph(str(ticket.get("description", "")))
    document.add_heading("CAPA Recommendation", level=2)
    document.add_paragraph(ticket_state.get("final_output") or ticket_state["draft"])
    document.add_heading("Relevant SOPs", level=2)
    for source in sorted({item["source"] for item in ticket_state.get("retrieved_context", [])}):
        document.add_paragraph(source, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


refresh_queue()

st.title("Quality Incident Queue")

if not workflow.excel_persistence_available:
    st.warning(
        "Excel persistence is unavailable in the Python interpreter running Streamlit. "
        "Install openpyxl there to update incident_tickets.xlsx."
    )

col1, col2, col3, col4, col5 = st.columns(5)
metrics = workflow.get_metrics(status_filter=st.session_state.get("status_filter", "non_approved"))
col1.metric("Total tickets", metrics["total_tickets"])
col2.metric("Open review items", metrics["open_tickets"])
col3.metric("Severity accuracy", f"{metrics.get('severity_accuracy', 0.0)}%")
col4.metric("Category accuracy", f"{metrics.get('category_accuracy', 0.0)}%")
col5.metric("Avg draft time", f"{metrics['avg_time_to_draft_sec']:.1f}s")
st.caption(f"Evaluated tickets: {metrics.get('tickets_processed', 0)}")

st.subheader("Incoming tickets")

with st.expander("Create a ticket and generate a CAPA"):
    with st.form("new_ticket_form"):
        new_product_line = st.text_input("Product line", value="General")
        new_description = st.text_area("Incident description", height=160)
        create_ticket = st.form_submit_button("Create ticket and generate CAPA")

    if create_ticket:
        if not new_description.strip():
            st.error("Enter an incident description before creating the ticket.")
        else:
            new_ticket = workflow.create_ticket(new_description, new_product_line)
            new_state = workflow.process_ticket(new_ticket)
            st.session_state["new_ticket_state"] = new_state
            st.session_state["selected_ticket_id"] = new_ticket["ticket_id"]
            st.success(f"Created {new_ticket['ticket_id']} and saved it to the CSV and Excel records.")

    new_ticket_state = st.session_state.get("new_ticket_state")
    if new_ticket_state:
        st.write(f"Generated CAPA for {new_ticket_state['ticket']['ticket_id']}")
        st.download_button(
            "Download CAPA document",
            data=build_capa_document(new_ticket_state),
            file_name=f"CAPA_{new_ticket_state['ticket']['ticket_id']}.{ 'docx' if Document else 'txt' }",
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                if Document
                else "text/plain"
            ),
            key="new_ticket_capa_download",
        )

filter_status, filter_severity = st.columns(2)
status_options = ["Non-approved only", "Approved only", "All tickets"]
status_filter_value = st.session_state.get("status_filter", "non_approved")
status_index = {"non_approved": 0, "approved": 1, "all": 2}.get(status_filter_value, 0)
with filter_status:
    status_option = st.selectbox("Ticket status", status_options, index=status_index)
with filter_severity:
    severity_option = st.selectbox(
        "Severity",
        ["All", "High", "Medium", "Low"],
        index=["All", "High", "Medium", "Low"].index(st.session_state.get("severity_filter", "All")),
    )

if status_option == "Non-approved only":
    st.session_state["status_filter"] = "non_approved"
elif status_option == "Approved only":
    st.session_state["status_filter"] = "approved"
else:
    st.session_state["status_filter"] = "all"

st.session_state["severity_filter"] = severity_option
refresh_queue()

queue_df = []
for item in st.session_state["tickets"]:
    queue_df.append({
        "ticket_id": item["ticket"]["ticket_id"],
        "product_line": item["ticket"].get("product_line", ""),
        "severity": item["severity"],
        "category": item["category"],
        "status": item["status"],
    })

if queue_df:
    selected_ticket_id = st.selectbox("Select ticket", [row["ticket_id"] for row in queue_df], index=0)
    st.session_state["selected_ticket_id"] = selected_ticket_id
    selected_ticket = get_ticket_by_id(selected_ticket_id)
else:
    selected_ticket = None
    st.info("No tickets match the current filter.")

if selected_ticket is not None:
    st.subheader(f"Ticket details: {selected_ticket['ticket']['ticket_id']}")
    st.write(selected_ticket["ticket"]["description"])
    st.write(f"Severity: {selected_ticket['severity']} | Category: {selected_ticket['category']} | Status: {selected_ticket['status']}")
    with st.expander("Relevant SOP excerpts"):
        for chunk in selected_ticket.get("retrieved_context", [])[:3]:
            st.markdown(f"**{chunk['source']}**")
            st.write(chunk["text"])

    st.subheader("CAPA draft")
    edited_draft = st.text_area("Review and edit the draft", value=selected_ticket["draft"], height=380)

    col_approve, col_reject = st.columns(2)
    with col_approve:
        if st.button("Approve draft"):
            selected_ticket = workflow.finalize_review(selected_ticket, "approve", edited_draft)
            selected_ticket["status"] = "approved"
            persist_ticket(selected_ticket)
            st.success("Draft approved and saved as approved in the source data.")
    with col_reject:
        if st.button("Reject draft"):
            selected_ticket = workflow.finalize_review(selected_ticket, "reject", edited_draft)
            selected_ticket["status"] = "rejected"
            persist_ticket(selected_ticket)
            st.warning("Draft rejected and returned for rework.")

    st.caption(f"Current status: {selected_ticket['status']}")

    if selected_ticket["status"] in {"resolved", "approved"}:
        st.write("Final output:")
        st.code(selected_ticket["final_output"], language="text")

    st.download_button(
        "Download CAPA document",
        data=build_capa_document(selected_ticket),
        file_name=f"CAPA_{selected_ticket['ticket']['ticket_id']}.{ 'docx' if Document else 'txt' }",
        mime=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if Document
            else "text/plain"
        ),
        key=f"capa_download_{selected_ticket['ticket']['ticket_id']}",
    )

